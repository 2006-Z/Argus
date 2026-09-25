import argparse
import os
import queue
import re
import shutil
import sys
import tempfile
import threading
import time

import openpyxl
import requests
from playwright.sync_api import sync_playwright

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from colors import BOLD, RESET, dim, fail, heading, ok, warn

# When stdout isn't attached to a real console, Python falls back to the
# OS default encoding (cp1252 on Windows) instead of UTF-8 -- printing
# anything outside that narrow charset then crashes with
# UnicodeEncodeError. errors="replace" makes that merely show a
# substitute character instead of taking the whole tool down.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

BASE_URL = "https://pcccvr.delhipolice.gov.in"
ELEMENT_ID_COL = "ELEMENT ID"
APP_NO_COL = "PCC APPLICATION NUMBER"
NAME_COL = "NAME"
FATHER_COL = "FATHER NAME"
INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')

PLAYWRIGHT_WORKERS = 2


def parse_args():
    parser = argparse.ArgumentParser(description="Check Delhi PCC application status; downloads certificates or screenshots rejections.")
    parser.add_argument("excel_file", help="Path to Excel file (ELEMENT ID, PCC APPLICATION NUMBER, NAME, FATHER NAME)")
    parser.add_argument("--workers", type=int, default=5, help="Parallel status checks (HTTP phase), max 20 (default: 5)")
    parser.add_argument("--out", default="delhi_pcc_status_checker_output", help="Output folder")
    return parser.parse_args()


def date_only_as_text(date_str: str) -> str:
    if not date_str:
        return ""
    date_part = date_str.split(",")[0].strip()
    return f"'{date_part}"


def sanitize_filename(name) -> str:
    name = str(name).strip()
    return INVALID_FILENAME_CHARS.sub("_", name)


def find_column_index(header_row, target: str):
    for idx, cell in enumerate(header_row):
        if cell is not None and str(cell).strip().lower() == target.lower():
            return idx
    return None


def parse_excel(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        sys.exit("Error: the Excel file is empty.")

    eid_idx = find_column_index(header, ELEMENT_ID_COL)
    app_idx = find_column_index(header, APP_NO_COL)
    name_idx = find_column_index(header, NAME_COL)
    father_idx = find_column_index(header, FATHER_COL)
    missing = []
    if eid_idx is None:
        missing.append(ELEMENT_ID_COL)
    if app_idx is None:
        missing.append(APP_NO_COL)
    if name_idx is None:
        missing.append(NAME_COL)
    if father_idx is None:
        missing.append(FATHER_COL)
    if missing:
        sys.exit(f"Error: missing required column(s): {', '.join(missing)}")

    rows = []
    for raw_row in rows_iter:
        if raw_row is None or all(v is None for v in raw_row):
            continue
        rows.append({
            "element_id": raw_row[eid_idx],
            "app_no": raw_row[app_idx],
            "name": raw_row[name_idx],
            "father": raw_row[father_idx],
        })
    return rows


def normalize_app_no(app_no: str) -> str:
    app_no = app_no.strip()
    if app_no.upper().startswith("DLSB-PCC/"):
        return app_no
    return f"DLSB-PCC/{app_no}"


ROW_COL = 6      # "999   "
ID_COL = 16      # element id
APP_COL_W = 16   # application number
TIME_COL = 7     # "12.34s"

# A cursor can only ever be moved up into rows that are STILL on screen --
# once a row scrolls past the top of the terminal, no escape code can
# address it again. So instead of one line per Excel row, there's a small
# FIXED block of "worker" lines that always sits at the bottom of the
# screen and is the only thing ever rewritten in place. Each row's final
# result is printed ONCE, as a permanent new line inserted just above
# that block -- never touched again. Phase 1 and phase 2 each set up
# their own slot block (different worker counts), one after the other.
_print_lock = threading.Lock()
_num_slots = 0
_slot_text = []


def init_slots(n, idle_line):
    global _num_slots, _slot_text
    _num_slots = n
    _slot_text = [idle_line] * n
    for line in _slot_text:
        print(line)


def update_slot(slot_idx, text):
    with _print_lock:
        _slot_text[slot_idx] = text
        lines_up = _num_slots - slot_idx
        print(f"\033[{lines_up}A\r\033[2K{text}\033[{lines_up}B\r", end="", flush=True)


def log_completed(line):
    with _print_lock:
        print(f"\033[{_num_slots}A\033[J", end="")
        print(line)
        for slot_line in _slot_text:
            print(slot_line)


def clear_slots():
    with _print_lock:
        print(f"\033[{_num_slots}A\033[J", end="", flush=True)


def run_ordered(indices, worker_count, idle_line, work_fn):
    """Runs `indices` (positions into whatever list) through a fresh
    worker-slot pool, in a fixed pool of `worker_count` threads that each
    keep the same slot index for their whole life. `work_fn(slot_idx, pos,
    i)` does the actual work and returns (result, line); results are
    logged in the SAME order as `indices` regardless of completion order.
    Returns {index: result}."""
    tasks = queue.Queue()
    for pos, i in enumerate(indices):
        tasks.put((pos, i))
    init_slots(worker_count, idle_line)

    results = {}
    results_lock = threading.Lock()
    pending = {}
    next_to_log = 0
    order_lock = threading.Lock()

    def submit(pos, line):
        nonlocal next_to_log
        with order_lock:
            pending[pos] = line
            while next_to_log in pending:
                entry = pending.pop(next_to_log)
                try:
                    log_completed(entry)
                except Exception:
                    pass  # never let one bad line jam every row after it
                next_to_log += 1

    def worker_loop(slot_idx):
        while True:
            try:
                pos, i = tasks.get_nowait()
            except queue.Empty:
                return
            result, line = work_fn(slot_idx, pos, i)
            with results_lock:
                results[i] = result
            submit(pos, line)

    threads = [threading.Thread(target=worker_loop, args=(w,)) for w in range(worker_count)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    clear_slots()
    return results


def color_bucket(bucket):
    if bucket == "SUCCESS":
        return ok(bucket)
    if bucket in ("REJECTED", "ERROR"):
        return fail(bucket)
    return warn(bucket)


def format_row1(row_num, element_id, app_no, time_text, status_text):
    row = str(row_num).ljust(ROW_COL)
    eid = str(element_id).ljust(ID_COL)[:ID_COL]
    app = str(app_no).ljust(APP_COL_W)[:APP_COL_W]
    time_ = time_text.rjust(TIME_COL)
    return f"{dim(row)}{eid}{app}{time_}  {status_text}"


IDLE_ROW1 = format_row1("", "", "", "", dim("..."))


def check_status(slot_idx, row_num, session, app_no_raw, name, father):
    t0 = time.time()
    app_no_full = normalize_app_no(str(app_no_raw))
    name = str(name).strip()
    father = str(father).strip()

    update_slot(slot_idx, format_row1(row_num, "", app_no_raw, "", warn("searching...")))
    try:
        r1 = session.post(
            f"{BASE_URL}/api/PccForm/search-pcc-application",
            json={"applicationNumber": app_no_full, "applicantName": name, "guardianName": father},
            timeout=20,
        )
        r1.raise_for_status()
        data1 = r1.json()
    except Exception as e:
        elapsed = time.time() - t0
        result = {"bucket": "ERROR", "status_text": "", "reason": f"Search failed: {e}", "date": "", "application_id": None}
        line = format_row1(row_num, "", app_no_raw, f"{elapsed:.2f}s", fail("ERROR") + f" - {result['reason'][:50]}")
        return result, line

    if not data1.get("success") or not data1.get("data"):
        elapsed = time.time() - t0
        reason = data1.get("message") or "No matching application found"
        result = {"bucket": "NOT_FOUND", "status_text": "", "reason": reason, "date": "", "application_id": None}
        line = format_row1(row_num, "", app_no_raw, f"{elapsed:.2f}s", warn("NOT_FOUND") + f" - {reason[:50]}")
        return result, line

    application_id = data1["data"][0]["applicationID"]

    update_slot(slot_idx, format_row1(row_num, "", app_no_raw, "", warn("checking status...")))
    try:
        r2 = session.post(
            f"{BASE_URL}/api/PccForm/get-pcc-application-status",
            json={"applicationId": str(application_id)},
            timeout=20,
        )
        r2.raise_for_status()
        data2 = r2.json()
    except Exception as e:
        elapsed = time.time() - t0
        result = {"bucket": "ERROR", "status_text": "", "reason": f"Status check failed: {e}", "date": "", "application_id": application_id}
        line = format_row1(row_num, "", app_no_raw, f"{elapsed:.2f}s", fail("ERROR") + f" - {result['reason'][:50]}")
        return result, line

    status_text = data2["data"]["status"]
    timeline = data2["data"].get("timeline") or []

    def last_date_for(stage_name):
        for entry in reversed(timeline):
            if entry.get("stage") == stage_name:
                return entry.get("date", "")
        return ""

    elapsed = time.time() - t0

    if status_text == "Certificate Issued":
        result = {"bucket": "SUCCESS", "status_text": status_text, "reason": "", "date": last_date_for("Certificate Issued"), "application_id": application_id}
        line = format_row1(row_num, "", app_no_raw, f"{elapsed:.2f}s", ok("SUCCESS"))
        return result, line

    if status_text == "Rejected":
        update_slot(slot_idx, format_row1(row_num, "", app_no_raw, "", warn("getting rejection reason...")))
        reason = ""
        try:
            r3 = session.post(f"{BASE_URL}/api/PccForm/get-pcc-applicant-details", json={"applicationID": str(application_id)}, timeout=20)
            r3.raise_for_status()
            data3 = r3.json()
            reason = (data3["data"]["applicantDetails"].get("rejectedRemarks") or "").strip()
        except Exception:
            pass
        if not reason:
            for entry in reversed(timeline):
                if entry.get("stage") == "Rejected":
                    reason = entry.get("description", "Rejected")
                    break
            else:
                reason = "Rejected"
        elapsed = time.time() - t0
        result = {"bucket": "REJECTED", "status_text": status_text, "reason": reason, "date": last_date_for("Rejected"), "application_id": application_id}
        line = format_row1(row_num, "", app_no_raw, f"{elapsed:.2f}s", fail("REJECTED") + f" - {reason[:50]}")
        return result, line

    result = {"bucket": "PENDING", "status_text": status_text, "reason": status_text, "date": timeline[-1]["date"] if timeline else "", "application_id": application_id}
    line = format_row1(row_num, "", app_no_raw, f"{elapsed:.2f}s", warn("PENDING") + f" - {status_text[:50]}")
    return result, line


def open_track_modal(page):
    page.goto(f"{BASE_URL}/login", wait_until="domcontentloaded", timeout=60000)
    page.get_by_text("PCC CONFIRMATION STATUS").click()
    page.get_by_placeholder("2026XXXXXXXX").wait_for(state="visible", timeout=20000)


def search_application(page, app_no, name, father):
    bare_app_no = str(app_no).replace("DLSB-PCC/", "").strip()
    app_input = page.get_by_placeholder("2026XXXXXXXX")
    app_input.click()
    app_input.fill(bare_app_no)
    name_input = page.get_by_placeholder("Enter applicant name")
    name_input.click()
    name_input.fill(str(name).strip())
    father_input = page.get_by_placeholder("Enter guardian name")
    father_input.click()
    father_input.fill(str(father).strip())
    page.get_by_role("button", name="Search Application").click()


_REJECTION_CLIP_JS = """
() => {
    function findByText(text) {
        const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_ELEMENT);
        let node;
        while (node = walker.nextNode()) {
            if (node.children.length === 0 && node.textContent.trim() === text) return node;
        }
        return null;
    }
    const appNumLabel = findByText('APPLICATION NUMBER');
    const redBox = document.querySelector('.alert-danger');
    const modalContent = document.querySelector('.modal-content');
    if (!appNumLabel || !redBox) return null;
    const topRect = appNumLabel.getBoundingClientRect();
    const redRect = redBox.getBoundingClientRect();
    const pad = 16;
    const left = Math.min(topRect.x, redRect.x) - pad / 2;
    const top = topRect.y - pad;
    const desiredRight = Math.max(topRect.right, redRect.right) + 20;
    const cardRight = modalContent ? modalContent.getBoundingClientRect().right - 4 : desiredRight;
    const right = Math.min(desiredRight, cardRight);
    return {
        x: left,
        y: top,
        width: right - left,
        height: (redRect.y + redRect.height) - top + pad / 2,
    };
}
"""


def capture_rejection_screenshot(page, dest_path):
    clip = page.evaluate(_REJECTION_CLIP_JS)
    if clip is None:
        page.screenshot(path=dest_path, full_page=True)
    else:
        page.screenshot(path=dest_path, clip=clip)


def wait_for_stable_rejection_text(page, timeout_ms=60000, stable_checks=3, check_interval_ms=400):
    page.get_by_text("Application Rejected").wait_for(state="visible", timeout=timeout_ms)
    deadline = time.time() + (timeout_ms / 1000)
    last_text = None
    stable_count = 0
    while time.time() < deadline:
        boxes = page.locator(".alert-danger")
        current_text = boxes.first.text_content() if boxes.count() > 0 else ""
        if current_text and "Application Rejected" in current_text and current_text == last_text:
            stable_count += 1
            if stable_count >= stable_checks:
                return
        else:
            stable_count = 0
        last_text = current_text
        page.wait_for_timeout(check_interval_ms)
    raise TimeoutError("Rejection content never stabilized")


def format_row2(row_num, element_id, time_text, status_text):
    row = str(row_num).ljust(ROW_COL)
    eid = str(element_id).ljust(ID_COL)[:ID_COL]
    time_ = time_text.rjust(TIME_COL)
    return f"{dim(row)}{eid}{time_}  {status_text}"


IDLE_ROW2 = format_row2("", "", "", dim("..."))


def process_phase2_entry(slot_idx, row_num, row, download_dir, screenshot_dir):
    t0 = time.time()
    element_id = sanitize_filename(row["element_id"])
    bucket = row["bucket"]

    update_slot(slot_idx, format_row2(row_num, element_id, "", warn("opening browser...")))
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(accept_downloads=True)
            page = context.new_page()
            try:
                open_track_modal(page)
                update_slot(slot_idx, format_row2(row_num, element_id, "", warn("searching...")))
                search_application(page, row["app_no"], row["name"], row["father"])

                if bucket == "SUCCESS":
                    update_slot(slot_idx, format_row2(row_num, element_id, "", warn("downloading certificate...")))
                    download_btn = page.get_by_role("button", name="Download Certificate")
                    download_btn.wait_for(state="visible", timeout=60000)
                    with page.expect_download(timeout=60000) as download_info:
                        download_btn.click()
                    download = download_info.value
                    dest = os.path.join(download_dir, f"{element_id}.pdf")
                    download.save_as(dest)
                    elapsed = time.time() - t0
                    result = {"phase2_status": "done", "output_file": dest}
                    line = format_row2(row_num, element_id, f"{elapsed:.2f}s", ok("done"))
                    return result, line
                else:
                    update_slot(slot_idx, format_row2(row_num, element_id, "", warn("waiting for rejection details...")))
                    wait_for_stable_rejection_text(page, timeout_ms=60000)
                    update_slot(slot_idx, format_row2(row_num, element_id, "", warn("capturing screenshot...")))
                    dest = os.path.join(screenshot_dir, f"{element_id}.png")
                    capture_rejection_screenshot(page, dest)
                    elapsed = time.time() - t0
                    result = {"phase2_status": "done", "output_file": dest}
                    line = format_row2(row_num, element_id, f"{elapsed:.2f}s", ok("done"))
                    return result, line
            finally:
                context.close()
                browser.close()
    except Exception as e:
        elapsed = time.time() - t0
        result = {"phase2_status": "failed", "output_file": "", "phase2_error": str(e)}
        line = format_row2(row_num, element_id, f"{elapsed:.2f}s", fail("failed") + f" - {str(e)[:50]}")
        return result, line


def format_duration(seconds):
    if seconds < 60:
        return f"{seconds:.1f}s"
    minutes, secs = divmod(seconds, 60)
    return f"{int(minutes)}m {secs:.0f}s"


def main():
    args = parse_args()
    if not os.path.exists(args.excel_file):
        sys.exit(f"Error: file not found: {args.excel_file}")

    workers = max(1, min(args.workers, 20))
    rows = parse_excel(args.excel_file)
    if not rows:
        sys.exit("Error: no data rows found in the Excel file.")

    total = len(rows)
    start_time = time.time()

    print(heading("Phase 1:") + f" checking status for {total} rows with {BOLD}{workers}{RESET} parallel HTTP requests...\n")
    header1 = (
        dim("Row".ljust(ROW_COL)) + "Element ID".ljust(ID_COL)
        + "Application".ljust(APP_COL_W) + "Time".rjust(TIME_COL) + "  Status"
    )
    print(header1)
    print(dim("-" * (ROW_COL + ID_COL + APP_COL_W + TIME_COL + 9)))

    def phase1_work(slot_idx, pos, i):
        session = requests.Session()
        session.headers.update({"Content-Type": "application/json", "Accept": "application/json"})
        return check_status(slot_idx, pos + 1, session, rows[i]["app_no"], rows[i]["name"], rows[i]["father"])

    phase1_results = run_ordered(list(range(total)), workers, IDLE_ROW1, phase1_work)

    phase2_indices = [i for i in range(total) if phase1_results[i]["bucket"] in ("SUCCESS", "REJECTED")]
    download_dir = tempfile.mkdtemp(prefix="pcc_downloads_")
    screenshot_dir = tempfile.mkdtemp(prefix="pcc_screenshots_")
    phase2_results = {}

    if phase2_indices:
        print(heading(f"\nPhase 2:") + f" downloading certificates / screenshotting rejections for {len(phase2_indices)} rows ({PLAYWRIGHT_WORKERS} parallel browsers)...\n")
        header2 = dim("Row".ljust(ROW_COL)) + "Element ID".ljust(ID_COL) + "Time".rjust(TIME_COL) + "  Status"
        print(header2)
        print(dim("-" * (ROW_COL + ID_COL + TIME_COL + 9)))

        def phase2_work(slot_idx, pos, i):
            row = dict(rows[i])
            row["bucket"] = phase1_results[i]["bucket"]
            return process_phase2_entry(slot_idx, pos + 1, row, download_dir, screenshot_dir)

        phase2_results = run_ordered(phase2_indices, PLAYWRIGHT_WORKERS, IDLE_ROW2, phase2_work)

    elapsed_total = time.time() - start_time

    counts = {}
    for r in phase1_results.values():
        counts[r["bucket"]] = counts.get(r["bucket"], 0) + 1
    counts_text = ", ".join(f"{color_bucket(b)}: {n}" for b, n in counts.items())
    print(heading("\nDone: ") + counts_text + f" Total time: {format_duration(elapsed_total)}.")

    report_wb = openpyxl.Workbook()
    ws = report_wb.active
    ws.title = "PCC Status Report"
    ws.append(["ELEMENT ID", "PCC APPLICATION NUMBER", "NAME", "FATHER NAME", "RESULT", "PORTAL STATUS", "REASON", "DATE", "OUTPUT FILE"])
    for i, row in enumerate(rows):
        r1 = phase1_results[i]
        r2 = phase2_results.get(i, {})
        reason = r1["reason"]
        if r2.get("phase2_status") == "failed":
            reason = (reason + " | " if reason else "") + f"Screenshot/download failed: {r2.get('phase2_error', '')}"
        ws.append([
            row["element_id"], row["app_no"], row["name"], row["father"],
            r1["bucket"], r1["status_text"], reason,
            date_only_as_text(r1.get("date", "")),
            os.path.basename(r2.get("output_file", "")) if r2.get("output_file") else "",
        ])

    os.makedirs(args.out, exist_ok=True)
    report_wb.save(os.path.join(args.out, "pcc_status_report.xlsx"))

    certs_dest = os.path.join(args.out, "certificates")
    os.makedirs(certs_dest, exist_ok=True)
    for fname in os.listdir(download_dir):
        shutil.move(os.path.join(download_dir, fname), os.path.join(certs_dest, fname))

    screenshots_dest = os.path.join(args.out, "rejected_screenshots")
    os.makedirs(screenshots_dest, exist_ok=True)
    for fname in os.listdir(screenshot_dir):
        shutil.move(os.path.join(screenshot_dir, fname), os.path.join(screenshots_dest, fname))

    shutil.rmtree(download_dir, ignore_errors=True)
    shutil.rmtree(screenshot_dir, ignore_errors=True)
    print(ok(f"Saved to folder: {args.out}"))


if __name__ == "__main__":
    main()
