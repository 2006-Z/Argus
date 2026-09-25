import argparse
import os
import queue
import re
import shutil
import sys
import tempfile
import threading
import time

import openpyxl
from playwright.sync_api import sync_playwright

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from colors import BOLD, RESET, dim, fail, heading, ok, warn

# When stdout isn't attached to a real console, Python falls back to the
# OS default encoding (cp1252 on Windows) instead of UTF-8 -- printing
# anything outside that narrow charset (e.g. Devanagari text scraped from
# a government site) then crashes with UnicodeEncodeError. errors="replace"
# makes that merely show a substitute character instead of taking the
# whole tool down.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

BASE_URL = "https://rajkosh.up.nic.in"
ELEMENT_ID_COL = "ELEMENT ID"
CHALLAN_COL = "CHALLAN"
INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def parse_args():
    parser = argparse.ArgumentParser(description="Search UP Challan numbers and download the certificate PDF for each.")
    parser.add_argument("excel_file", help="Path to the Excel file (must have ELEMENT ID and CHALLAN columns)")
    parser.add_argument("--workers", type=int, default=2, help="Parallel browser tasks, max 5 (default: 2)")
    parser.add_argument("--out", default="up_challan_downloader_output", help="Output folder (default: up_challan_downloader_output)")
    return parser.parse_args()


def sanitize_filename(name) -> str:
    name = str(name).strip()
    return INVALID_FILENAME_CHARS.sub("_", name)


def find_column_index(header_row, target: str):
    for idx, cell in enumerate(header_row):
        if cell is not None and str(cell).strip().lower() == target.lower():
            return idx
    return None


def parse_excel(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        sys.exit("Error: the Excel file is empty.")

    eid_idx = find_column_index(header, ELEMENT_ID_COL)
    challan_idx = find_column_index(header, CHALLAN_COL)
    missing = []
    if eid_idx is None:
        missing.append(ELEMENT_ID_COL)
    if challan_idx is None:
        missing.append(CHALLAN_COL)
    if missing:
        sys.exit(f"Error: missing required column(s): {', '.join(missing)}")

    rows = []
    for raw_row in rows_iter:
        if raw_row is None or all(v is None for v in raw_row):
            continue
        rows.append({
            "element_id": raw_row[eid_idx],
            "challan": raw_row[challan_idx],
        })
    return rows


_HIDE_PRINT_CLOSE_LINKS_JS = """
() => {
    const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_ELEMENT);
    let node;
    while (node = walker.nextNode()) {
        if (node.tagName === 'A' && (node.textContent.trim() === 'Print' || node.textContent.trim() === 'Close')) {
            node.style.display = 'none';
        }
    }
}
"""


def open_search_page(page):
    page.goto(f"{BASE_URL}/default.aspx", wait_until="domcontentloaded", timeout=60000)
    page.get_by_role("link", name="Search and Print Challan").click()
    page.locator("#txtChallanNo").wait_for(state="visible", timeout=20000)


def search_challan(page, challan_no):
    challan_input = page.locator("#txtChallanNo")
    challan_input.click()
    challan_input.fill(str(challan_no).strip())
    page.locator("#btnSearch").click()


def wait_for_search_result(page, timeout_ms=60000):
    deadline = time.time() + (timeout_ms / 1000)
    while time.time() < deadline:
        if page.get_by_text("No. Record Found.").count() > 0:
            return "not_found"
        if page.get_by_text("Click For Detail").count() > 0:
            return "found"
        page.wait_for_timeout(400)
    raise TimeoutError("Search result never appeared")


ROW_COL = 6     # "999   "
ID_COL = 16     # element id
CHALLAN_COL_W = 14  # challan number
TIME_COL = 7    # "12.34s"

# A cursor can only ever be moved up into rows that are STILL on screen --
# once a row scrolls past the top of the terminal, no escape code can
# address it again. So instead of one line per Excel row, there's a small
# FIXED block of "worker" lines that always sits at the bottom of the
# screen and is the only thing ever rewritten in place. Each row's final
# result is printed ONCE, as a permanent new line inserted just above
# that block -- never touched again.
_print_lock = threading.Lock()
_num_slots = 0
_slot_text = []


def format_row(row_num, element_id, challan_text, time_text, status_text):
    row = str(row_num).ljust(ROW_COL)
    eid = str(element_id).ljust(ID_COL)[:ID_COL]
    challan = str(challan_text).ljust(CHALLAN_COL_W)[:CHALLAN_COL_W]
    time_ = time_text.rjust(TIME_COL)
    return f"{dim(row)}{eid}{challan}{time_}  {status_text}"


IDLE_SLOT = format_row("", "", "", "", dim("..."))


def init_slots(n):
    global _num_slots, _slot_text
    _num_slots = n
    _slot_text = [IDLE_SLOT] * n
    for line in _slot_text:
        print(line)


def update_slot(slot_idx, text):
    with _print_lock:
        _slot_text[slot_idx] = text
        lines_up = _num_slots - slot_idx
        print(f"\033[{lines_up}A\r\033[2K{text}\033[{lines_up}B\r", end="", flush=True)


def log_completed(line):
    with _print_lock:
        print(f"\033[{_num_slots}A\033[J", end="")
        print(line)
        for slot_line in _slot_text:
            print(slot_line)


def clear_slots():
    with _print_lock:
        print(f"\033[{_num_slots}A\033[J", end="", flush=True)


def color_bucket(bucket):
    if bucket == "SUCCESS":
        return ok(bucket)
    if bucket in ("FAILURE", "ERROR"):
        return fail(bucket)
    return warn(bucket)


def process_challan(slot_idx, row_num, row, pdf_dir):
    t0 = time.time()
    element_id = sanitize_filename(row["element_id"])
    challan_no = row["challan"]
    challan_text = "" if challan_no is None else str(challan_no).strip()

    if challan_text == "":
        reason = "CHALLAN number missing"
        line = format_row(row_num, element_id, "", "0.00s", fail("NOT_FOUND") + f" - {reason}")
        return "NOT_FOUND", reason, "", line

    update_slot(slot_idx, format_row(row_num, element_id, challan_text, "", warn("opening browser...")))
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(accept_downloads=True)
            page = context.new_page()
            try:
                open_search_page(page)
                update_slot(slot_idx, format_row(row_num, element_id, challan_text, "", warn("searching...")))
                search_challan(page, challan_no)
                update_slot(slot_idx, format_row(row_num, element_id, challan_text, "", warn("waiting for result...")))
                outcome = wait_for_search_result(page, timeout_ms=60000)

                if outcome == "not_found":
                    elapsed = time.time() - t0
                    reason = "No record found"
                    line = format_row(row_num, element_id, challan_text, f"{elapsed:.2f}s", warn("NOT_FOUND") + f" - {reason}")
                    return "NOT_FOUND", reason, "", line

                table = page.locator("#GridViewChallan")
                data_row = table.locator("tr").nth(1)
                row_text = " ".join(data_row.inner_text().split())

                if "Success" in row_text:
                    update_slot(slot_idx, format_row(row_num, element_id, challan_text, "", warn("downloading PDF...")))
                    print_btn = data_row.get_by_role("button", name="Print Challan")
                    with page.context.expect_page() as popup_info:
                        print_btn.click()
                    popup = popup_info.value
                    popup.wait_for_load_state("domcontentloaded", timeout=30000)
                    popup.get_by_text("CHALLAN", exact=True).first.wait_for(state="visible", timeout=30000)
                    popup.wait_for_timeout(500)
                    popup.evaluate(_HIDE_PRINT_CLOSE_LINKS_JS)
                    dest = os.path.join(pdf_dir, f"{element_id}.pdf")
                    popup.pdf(path=dest, print_background=True, margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
                    popup.close()
                    elapsed = time.time() - t0
                    line = format_row(row_num, element_id, challan_text, f"{elapsed:.2f}s", ok("SUCCESS") + f" - {row_text[:60]}")
                    return "SUCCESS", row_text, dest, line

                if "Failure" in row_text:
                    elapsed = time.time() - t0
                    line = format_row(row_num, element_id, challan_text, f"{elapsed:.2f}s", fail("FAILURE") + f" - {row_text[:60]}")
                    return "FAILURE", row_text, "", line

                elapsed = time.time() - t0
                line = format_row(row_num, element_id, challan_text, f"{elapsed:.2f}s", warn("UNKNOWN") + f" - {row_text[:60]}")
                return "UNKNOWN", row_text, "", line
            finally:
                context.close()
                browser.close()
    except Exception as e:
        elapsed = time.time() - t0
        reason = str(e)
        line = format_row(row_num, element_id, challan_text, f"{elapsed:.2f}s", fail("ERROR") + f" - {reason[:60]}")
        return "ERROR", reason, "", line


def format_duration(seconds):
    if seconds < 60:
        return f"{seconds:.1f}s"
    minutes, secs = divmod(seconds, 60)
    return f"{int(minutes)}m {secs:.0f}s"


def run_pass(indices, rows, pdf_dir, workers):
    """Runs `indices` (positions into `rows`) through the worker-slot
    pool, logging each result in the SAME order as `indices` itself, and
    returns {index: (bucket, reason, output_file)}."""
    tasks = queue.Queue()
    for pos, i in enumerate(indices):
        tasks.put((pos, i))
    init_slots(workers)

    results = {}
    results_lock = threading.Lock()
    pending = {}
    next_to_log = 0
    order_lock = threading.Lock()

    def submit_result(pos, line):
        nonlocal next_to_log
        with order_lock:
            pending[pos] = line
            while next_to_log in pending:
                entry = pending.pop(next_to_log)
                try:
                    log_completed(entry)
                except Exception:
                    pass  # never let one bad line jam every row after it
                next_to_log += 1

    def worker_loop(slot_idx):
        while True:
            try:
                pos, i = tasks.get_nowait()
            except queue.Empty:
                return
            bucket, reason, output_file, line = process_challan(slot_idx, pos + 1, rows[i], pdf_dir)
            with results_lock:
                results[i] = (bucket, reason, output_file)
            submit_result(pos, line)

    threads = [threading.Thread(target=worker_loop, args=(w,)) for w in range(workers)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    clear_slots()
    return results


def main():
    args = parse_args()
    if not os.path.exists(args.excel_file):
        sys.exit(f"Error: file not found: {args.excel_file}")

    workers = max(1, min(args.workers, 5))
    rows = parse_excel(args.excel_file)
    if not rows:
        sys.exit("Error: no data rows found in the Excel file.")

    total = len(rows)
    print(heading(f"Found {total} rows.") + f" Checking with {BOLD}{workers}{RESET} parallel browser tasks...\n")

    header = (
        dim("Row".ljust(ROW_COL)) + "Element ID".ljust(ID_COL)
        + "Challan".ljust(CHALLAN_COL_W) + "Time".rjust(TIME_COL) + "  Status"
    )
    print(header)
    print(dim("-" * (ROW_COL + ID_COL + CHALLAN_COL_W + TIME_COL + 9)))

    pdf_dir = tempfile.mkdtemp(prefix="challan_pdfs_")

    start_time = time.time()
    results = run_pass(list(range(total)), rows, pdf_dir, workers)

    # One retry pass for ERROR rows only (our own automation glitch, not a
    # real site answer) -- printed as its own fresh block below the main
    # results, since the original rows' lines are already permanent.
    retry_indices = [i for i in range(total) if results[i][0] == "ERROR"]
    if retry_indices:
        print(warn(f"\nRetrying {len(retry_indices)} row(s) that hit an automation error..."))
        retry_results = run_pass(retry_indices, rows, pdf_dir, workers)
        results.update(retry_results)
    elapsed_total = time.time() - start_time

    counts = {}
    for i in range(total):
        bucket = results[i][0]
        counts[bucket] = counts.get(bucket, 0) + 1
    counts_text = ", ".join(f"{color_bucket(b)}: {n}" for b, n in counts.items())
    print(heading("\nDone: ") + counts_text + f" Total time: {format_duration(elapsed_total)}.")

    report_wb = openpyxl.Workbook()
    ws = report_wb.active
    ws.title = "Challan Report"
    ws.append(["ELEMENT ID", "CHALLAN", "RESULT", "DETAIL", "OUTPUT FILE"])
    for i, row in enumerate(rows):
        bucket, reason, output_file = results[i]
        ws.append([row["element_id"], row["challan"], bucket, reason, os.path.basename(output_file) if output_file else ""])

    os.makedirs(args.out, exist_ok=True)
    report_wb.save(os.path.join(args.out, "challan_report.xlsx"))

    challans_dir = os.path.join(args.out, "challans")
    os.makedirs(challans_dir, exist_ok=True)
    for fname in os.listdir(pdf_dir):
        shutil.move(os.path.join(pdf_dir, fname), os.path.join(challans_dir, fname))

    shutil.rmtree(pdf_dir, ignore_errors=True)
    print(ok(f"Saved to folder: {args.out}"))


if __name__ == "__main__":
    main()
