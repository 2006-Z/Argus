import argparse
import os
import queue
import re
import shutil
import sys
import tempfile
import threading
import time

import openpyxl
from playwright.sync_api import sync_playwright

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from colors import BOLD, RESET, dim, fail, heading, ok, warn

# When stdout isn't attached to a real console, Python falls back to the
# OS default encoding (cp1252 on Windows) instead of UTF-8 -- printing
# anything outside that narrow charset (e.g. Devanagari text scraped from
# a government site) then crashes with UnicodeEncodeError. errors="replace"
# makes that merely show a substitute character instead of taking the
# whole tool down.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

BASE_URL = "https://cctnsup.gov.in"
ELEMENT_ID_COL = "ELEMENT ID"
APPLICATION_COL = "APPLICATION"
INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')

REJECTED_MARKER = "अस्वीकृत"
ACCEPTED_MARKER = "स्वीकृत"
NOT_FOUND_MARKER = "No Data Found"

STATUS_CELL_INDEX = 1
DATE_CELL_INDEX = 4


def parse_args():
    parser = argparse.ArgumentParser(description="Check UP CCTNS application status; screenshots rejected cases.")
    parser.add_argument("excel_file", help="Path to the Excel file (must have ELEMENT ID and APPLICATION columns)")
    parser.add_argument("--workers", type=int, default=2, help="Parallel browser tasks, max 5 (default: 2)")
    parser.add_argument("--out", default="up_pcc_status_checker_output", help="Output folder")
    return parser.parse_args()


def sanitize_filename(name) -> str:
    name = str(name).strip()
    return INVALID_FILENAME_CHARS.sub("_", name)


def find_column_index(header_row, target: str):
    for idx, cell in enumerate(header_row):
        if cell is not None and str(cell).strip().lower() == target.lower():
            return idx
    return None


def parse_excel(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        sys.exit("Error: the Excel file is empty.")

    eid_idx = find_column_index(header, ELEMENT_ID_COL)
    app_idx = find_column_index(header, APPLICATION_COL)
    missing = []
    if eid_idx is None:
        missing.append(ELEMENT_ID_COL)
    if app_idx is None:
        missing.append(APPLICATION_COL)
    if missing:
        sys.exit(f"Error: missing required column(s): {', '.join(missing)}")

    rows = []
    for raw_row in rows_iter:
        if raw_row is None or all(v is None for v in raw_row):
            continue
        rows.append({
            "element_id": raw_row[eid_idx],
            "application": raw_row[app_idx],
        })
    return rows


_RESULT_CLIP_JS = """
() => {
    const table = document.getElementById('gdvdata');
    if (!table) return null;
    const tableRect = table.getBoundingClientRect();
    const pad = 16;
    return {
        x: 0,
        y: 0,
        width: Math.max(document.documentElement.clientWidth, tableRect.right + pad),
        height: tableRect.y + tableRect.height + pad,
    };
}
"""


def wait_for_stable_screenshot_region(page, timeout_ms=15000, stable_checks=3, check_interval_ms=300):
    deadline = time.time() + (timeout_ms / 1000)
    last_height = None
    stable_count = 0
    while time.time() < deadline:
        clip = page.evaluate(_RESULT_CLIP_JS)
        height = clip["height"] if clip else None
        if height is not None and height == last_height:
            stable_count += 1
            if stable_count >= stable_checks:
                return
        else:
            stable_count = 0
        last_height = height
        page.wait_for_timeout(check_interval_ms)


def capture_result_screenshot(page, dest_path):
    wait_for_stable_screenshot_region(page)
    clip = page.evaluate(_RESULT_CLIP_JS)
    if clip is None:
        page.screenshot(path=dest_path, full_page=True)
    else:
        page.screenshot(path=dest_path, clip=clip)


def wait_for_matching_result(page, app_no, timeout_ms=60000):
    deadline = time.time() + (timeout_ms / 1000)
    app_no_str = str(app_no).strip()
    while time.time() < deadline:
        table = page.locator("#gdvdata")
        if table.count() > 0:
            data_row = table.locator("tr").nth(1)
            try:
                row_text = data_row.inner_text()
            except Exception:
                row_text = ""
            if NOT_FOUND_MARKER in row_text:
                return "not_found", row_text, None
            if app_no_str and app_no_str in row_text:
                cells = data_row.locator("td").all_inner_texts()
                return "found", row_text, cells
        page.wait_for_timeout(400)
    raise TimeoutError(f"Result for {app_no_str} never appeared/matched")


ROW_COL = 6      # "999   "
ID_COL = 16      # element id
APP_COL_W = 16   # application number
TIME_COL = 7     # "12.34s"

# A cursor can only ever be moved up into rows that are STILL on screen --
# once a row scrolls past the top of the terminal, no escape code can
# address it again. So instead of one line per Excel row, there's a small
# FIXED block of "worker" lines that always sits at the bottom of the
# screen and is the only thing ever rewritten in place. Each row's final
# result is printed ONCE, as a permanent new line inserted just above
# that block -- never touched again.
_print_lock = threading.Lock()
_num_slots = 0
_slot_text = []


def format_row(row_num, element_id, app_text, time_text, status_text):
    row = str(row_num).ljust(ROW_COL)
    eid = str(element_id).ljust(ID_COL)[:ID_COL]
    app = str(app_text).ljust(APP_COL_W)[:APP_COL_W]
    time_ = time_text.rjust(TIME_COL)
    return f"{dim(row)}{eid}{app}{time_}  {status_text}"


IDLE_SLOT = format_row("", "", "", "", dim("..."))


def init_slots(n):
    global _num_slots, _slot_text
    _num_slots = n
    _slot_text = [IDLE_SLOT] * n
    for line in _slot_text:
        print(line)


def update_slot(slot_idx, text):
    with _print_lock:
        _slot_text[slot_idx] = text
        lines_up = _num_slots - slot_idx
        print(f"\033[{lines_up}A\r\033[2K{text}\033[{lines_up}B\r", end="", flush=True)


def log_completed(line):
    with _print_lock:
        print(f"\033[{_num_slots}A\033[J", end="")
        print(line)
        for slot_line in _slot_text:
            print(slot_line)


def clear_slots():
    with _print_lock:
        print(f"\033[{_num_slots}A\033[J", end="", flush=True)


def color_bucket(bucket):
    if bucket == "ACCEPTED":
        return ok(bucket)
    if bucket in ("REJECTED", "ERROR"):
        return fail(bucket)
    return warn(bucket)


def process_application(slot_idx, row_num, row, screenshot_dir):
    t0 = time.time()
    element_id = sanitize_filename(row["element_id"])
    app_no_raw = row["application"]
    if app_no_raw is None or str(app_no_raw).strip() == "":
        reason = "Application number missing"
        line = format_row(row_num, element_id, "", "0.00s", warn("NOT_FOUND") + f" - {reason}")
        return "NOT_FOUND", reason, "", "", line

    app_no = str(app_no_raw).strip().lstrip("`").strip()

    update_slot(slot_idx, format_row(row_num, element_id, app_no, "", warn("opening browser...")))
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context()
            page = context.new_page()
            try:
                page.goto(f"{BASE_URL}/citizenportal/CitizenVerification.aspx", wait_until="domcontentloaded", timeout=60000)
                page.locator("#txtServiceNo").wait_for(state="visible", timeout=20000)
                page.locator("#txtServiceNo").fill(app_no)
                update_slot(slot_idx, format_row(row_num, element_id, app_no, "", warn("searching...")))
                page.locator("#btnSearchFir").click()

                update_slot(slot_idx, format_row(row_num, element_id, app_no, "", warn("waiting for result...")))
                outcome, row_text, cells = wait_for_matching_result(page, app_no, timeout_ms=60000)
                if outcome == "not_found":
                    elapsed = time.time() - t0
                    reason = "No record found"
                    line = format_row(row_num, element_id, app_no, f"{elapsed:.2f}s", warn("NOT_FOUND") + f" - {reason}")
                    return "NOT_FOUND", reason, "", "", line

                row_text_clean = " ".join(row_text.split())
                date = cells[DATE_CELL_INDEX].strip() if cells and len(cells) > DATE_CELL_INDEX else ""
                status_cell_text = cells[STATUS_CELL_INDEX].strip() if cells and len(cells) > STATUS_CELL_INDEX else row_text_clean

                if REJECTED_MARKER in row_text_clean:
                    update_slot(slot_idx, format_row(row_num, element_id, app_no, "", warn("capturing screenshot...")))
                    dest = os.path.join(screenshot_dir, f"{element_id}.png")
                    capture_result_screenshot(page, dest)
                    elapsed = time.time() - t0
                    line = format_row(row_num, element_id, app_no, f"{elapsed:.2f}s", fail("REJECTED") + f" - {status_cell_text[:50]}")
                    return "REJECTED", status_cell_text, date, dest, line

                if ACCEPTED_MARKER in row_text_clean:
                    elapsed = time.time() - t0
                    line = format_row(row_num, element_id, app_no, f"{elapsed:.2f}s", ok("ACCEPTED"))
                    return "ACCEPTED", "Accepted", date, "", line

                elapsed = time.time() - t0
                line = format_row(row_num, element_id, app_no, f"{elapsed:.2f}s", warn("UNKNOWN") + f" - {status_cell_text[:50]}")
                return "UNKNOWN", status_cell_text, date, "", line
            finally:
                context.close()
                browser.close()
    except Exception as e:
        elapsed = time.time() - t0
        reason = str(e)
        line = format_row(row_num, element_id, app_no, f"{elapsed:.2f}s", fail("ERROR") + f" - {reason[:50]}")
        return "ERROR", reason, "", "", line


def format_duration(seconds):
    if seconds < 60:
        return f"{seconds:.1f}s"
    minutes, secs = divmod(seconds, 60)
    return f"{int(minutes)}m {secs:.0f}s"


def run_pass(indices, rows, screenshot_dir, workers):
    """Runs `indices` (positions into `rows`) through the worker-slot
    pool, logging each result in the SAME order as `indices` itself, and
    returns {index: (bucket, reason, date, output_file)}."""
    tasks = queue.Queue()
    for pos, i in enumerate(indices):
        tasks.put((pos, i))
    init_slots(workers)

    results = {}
    results_lock = threading.Lock()
    pending = {}
    next_to_log = 0
    order_lock = threading.Lock()

    def submit_result(pos, line):
        nonlocal next_to_log
        with order_lock:
            pending[pos] = line
            while next_to_log in pending:
                entry = pending.pop(next_to_log)
                try:
                    log_completed(entry)
                except Exception:
                    pass  # never let one bad line jam every row after it
                next_to_log += 1

    def worker_loop(slot_idx):
        while True:
            try:
                pos, i = tasks.get_nowait()
            except queue.Empty:
                return
            bucket, reason, date, output_file, line = process_application(slot_idx, pos + 1, rows[i], screenshot_dir)
            with results_lock:
                results[i] = (bucket, reason, date, output_file)
            submit_result(pos, line)

    threads = [threading.Thread(target=worker_loop, args=(w,)) for w in range(workers)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    clear_slots()
    return results


def main():
    args = parse_args()
    if not os.path.exists(args.excel_file):
        sys.exit(f"Error: file not found: {args.excel_file}")

    workers = max(1, min(args.workers, 5))
    rows = parse_excel(args.excel_file)
    if not rows:
        sys.exit("Error: no data rows found in the Excel file.")

    total = len(rows)
    print(heading(f"Found {total} rows.") + f" Checking with {BOLD}{workers}{RESET} parallel browser tasks...\n")

    header = (
        dim("Row".ljust(ROW_COL)) + "Element ID".ljust(ID_COL)
        + "Application".ljust(APP_COL_W) + "Time".rjust(TIME_COL) + "  Status"
    )
    print(header)
    print(dim("-" * (ROW_COL + ID_COL + APP_COL_W + TIME_COL + 9)))

    screenshot_dir = tempfile.mkdtemp(prefix="cctns_screenshots_")

    start_time = time.time()
    results = run_pass(list(range(total)), rows, screenshot_dir, workers)

    retry_indices = [i for i in range(total) if results[i][0] == "ERROR"]
    if retry_indices:
        print(warn(f"\nRetrying {len(retry_indices)} row(s) that hit an automation error..."))
        retry_results = run_pass(retry_indices, rows, screenshot_dir, workers)
        results.update(retry_results)
    elapsed_total = time.time() - start_time

    counts = {}
    for i in range(total):
        bucket = results[i][0]
        counts[bucket] = counts.get(bucket, 0) + 1
    counts_text = ", ".join(f"{color_bucket(b)}: {n}" for b, n in counts.items())
    print(heading("\nDone: ") + counts_text + f" Total time: {format_duration(elapsed_total)}.")

    report_wb = openpyxl.Workbook()
    ws = report_wb.active
    ws.title = "Application Status"
    ws.append(["ELEMENT ID", "APPLICATION", "RESULT", "APPLICATION DATE", "DETAIL", "OUTPUT FILE"])
    for i, row in enumerate(rows):
        bucket, reason, date, output_file = results[i]
        ws.append([
            row["element_id"],
            row["application"],
            bucket,
            f"'{date}" if date else "",
            reason,
            os.path.basename(output_file) if output_file else "",
        ])

    os.makedirs(args.out, exist_ok=True)
    report_wb.save(os.path.join(args.out, "application_status_report.xlsx"))

    screenshots_dest = os.path.join(args.out, "rejected_screenshots")
    os.makedirs(screenshots_dest, exist_ok=True)
    for fname in os.listdir(screenshot_dir):
        shutil.move(os.path.join(screenshot_dir, fname), os.path.join(screenshots_dest, fname))

    shutil.rmtree(screenshot_dir, ignore_errors=True)
    print(ok(f"Saved to folder: {args.out}"))


if __name__ == "__main__":
    main()
